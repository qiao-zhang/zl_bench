import sys

import pytest

import ab_compare
import rlc_ops
import rlc_report


TABLE = """\
RLC Enhance 开启前后算子性能对比,,,,,,
环境,,,,,,
,,,,,,
#,算子,bf16,,,,,
,,收益,,,退化,,
,,最弱,最强,shape,最轻,最重,shape,
1,kron,,,,,,
2,"outer
vllm:outer",,,,,,
3,"uniform_
main:uniform_inplace",,,,,,
"""


def _table(tmp_path):
    path = tmp_path / "table.csv"
    path.write_text(TABLE, encoding="utf-8")
    return path


def test_parse_cell_reads_the_display_name_and_inventory_op():
    assert rlc_ops.parse_cell("kron") == ("kron", "", "kron", None)
    assert rlc_ops.parse_cell("outer\nvllm:outer") == ("outer", "vllm", "outer", None)
    # the table's label and the inventory op differ for two tests that share a label
    assert rlc_ops.parse_cell("uniform_\nmain:uniform_inplace") == \
        ("uniform_", "main", "uniform_inplace", None)
    assert rlc_ops.parse_cell("x\nmain:y@z") == ("x", "main", "y", "z")
    assert rlc_ops.parse_cell("vllm:foo") == ("vllm:foo", "vllm", "foo", None)
    # a single line carries the same `[repo:]op[@op_file]` spec as the second line
    assert rlc_ops.parse_cell("main:kron@kron") == ("main:kron@kron", "main", "kron", "kron")
    assert rlc_ops.parse_cell("kron@kron") == ("kron@kron", "", "kron", "kron")
    assert rlc_ops.parse_cell("") is None


def test_read_table_skips_the_title_and_header_rows(tmp_path):
    entries = rlc_ops.read_table(_table(tmp_path))
    assert [(e.display, e.repo, e.op) for e in entries] == [
        ("kron", "", "kron"),
        ("outer", "vllm", "outer"),
        ("uniform_", "main", "uniform_inplace"),
    ]


def test_resolve_picks_the_only_op_file(monkeypatch, tmp_path):
    inv = {
        "main": [{"op": "kron", "op_file": "kron", "collected": "yes"},
                 {"op": "uniform_inplace", "op_file": "uniform", "collected": "yes"}],
        "vllm": [{"op": "outer", "op_file": "outer", "collected": "yes"}],
    }
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(_table(tmp_path)))
    assert specs == ["main:kron@kron", "vllm:outer@outer", "main:uniform_inplace@uniform"]
    assert not errors


def test_resolve_expands_every_test_for_one_name(monkeypatch, tmp_path):
    path = tmp_path / "one.csv"
    path.write_text("#,算子\n1,kron\n", encoding="utf-8")
    inv = {"main": [{"op": "kron", "op_file": "kron", "collected": "yes"},
                    {"op": "kron", "op_file": "kron_other", "collected": "yes"}]}
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(path))
    assert specs == ["main:kron@kron", "main:kron@kron_other"]
    assert [r["status"] for r in report] == ["multi", "multi"]
    assert not errors


def _row(label="main:kron@kron", dtype="torch.bfloat16", gap=0.5, a=0.05, b=0.02,
         gate="changed", flags="", verdict="faster", row=0, confirm=""):
    return rlc_report.Row(label, "main/kron", "kron", dtype, row, gap, a, b,
                          gate, flags, verdict, confirm)


def test_summarizing_splits_gain_and_loss_rows():
    rows = [_row(gap=1.5, row=0), _row(gap=0.5, row=1),
            _row(gap=-0.2, row=2, a=0.02, b=0.025)]
    stats = rlc_report.summarizing(rows, {}, None, {}, "main:kron@kron", "bf16")
    assert stats["rows"] == 3
    assert stats["gain"] == (0.5, 1.5, 2)
    assert stats["loss"] == (-0.2, -0.2, 1)
    assert stats["cvt"] == (None, None, False)


def test_read_ab_csv_carries_the_confirm_column(tmp_path):
    """The recheck outcome has to reach the table: `timer_mismatch` there is the only
    place the report can say a verdict was never confirmed."""
    path = tmp_path / "rows.csv"
    path.write_text(
        "op,run,op_name,dtype,row,gap,median_a_ms,median_b_ms,gate,flags,verdict,confirm\n"
        "main:kron@kron,main/kron,kron,torch.bfloat16,0,0.5,0.05,0.02,changed,,faster,"
        "timer_mismatch\n")
    row = rlc_report.read_ab_csv(path)[0]
    assert row.confirm == "timer_mismatch"
    assert row.flags == "" and row.gap == 0.5


def test_read_ab_csv_tolerates_a_csv_without_the_confirm_column(tmp_path):
    """A --csv written before this column existed still loads (confirm defaults to "")."""
    path = tmp_path / "rows.csv"
    path.write_text(
        "op,run,op_name,dtype,row,gap,median_a_ms,median_b_ms,gate,flags,verdict\n"
        "main:kron@kron,main/kron,kron,torch.bfloat16,0,0.5,0.05,0.02,changed,,faster\n")
    assert rlc_report.read_ab_csv(path)[0].confirm == ""


def test_row_notes_separate_a_mismatched_main_round_from_a_refused_recheck():
    """Two warnings, not one: a mismatched main round makes that row's gap uncomparable,
    while a refused recheck only leaves the main-round verdict unconfirmed — the main
    rounds themselves are clean, so they must not be called uncomparable."""
    mismatched = [_row(flags="timer_mismatch", gap=0.2)]
    assert rlc_report.row_notes(mismatched, None, {}, "main:kron@kron") == [
        "A/B 计时口径不一致，性能差异不可直接比较（未计入收益/损失汇总，1 行）"]
    refused = [_row(gap=0.2, confirm="timer_mismatch"),
               _row(row=1, gap=0.1, confirm="timer_mismatch")]
    assert rlc_report.row_notes(refused, None, {}, "main:kron@kron") == [
        "复测计时口径不一致，主轮次结论未确认（2 行，未计入汇总）"]
    # --include-unconfirmed: the note must say the numbers contain them after all
    assert rlc_report.row_notes(mismatched, None, {}, "main:kron@kron", dropping=False) == [
        "A/B 计时口径不一致，性能差异不可直接比较（仍计入汇总）"]


def test_summarizing_excludes_unconfirmed_rows_by_default():
    """Rows whose timing is not comparable stay counted in `rows` (the table still warns
    about them) but drop out of the gain/loss ranges, and the note reports how many."""
    rows = [_row(gap=0.5, flags="timer_mismatch"),
            _row(row=1, gap=-0.2, a=0.02, b=0.025, confirm="timer_mismatch"),
            _row(row=2, gap=0.3)]
    stats = rlc_report.summarizing(rows, {}, None, {}, "main:kron@kron", "bf16")
    assert stats["rows"] == 3                      # 仍计入行数（诊断展示）
    assert stats["gain"] == (0.3, 0.3, 1)          # 只有干净行进收益区间
    assert stats["loss"] == (None, None, 0)
    assert "未计入收益/损失汇总，1 行" in stats["note"]
    assert "1 行，未计入汇总" in stats["note"]        # 复测被拒的那 1 行


def test_unconfirmed_covers_every_confirm_state_but_a_plain_undecided():
    """Every confirm state means the recheck settled nothing or refuted the verdict, so
    none may count as a measured gain. A plain undecided row (no confirm) measured
    nothing against it yet and stays in the summaries."""
    for state in ("timer_mismatch", "insufficient_rounds", "not_reproduced"):
        assert rlc_report.unconfirmed(_row(confirm=state))
    assert rlc_report.unconfirmed(_row(flags="timer_mismatch"))
    assert rlc_report.unconfirmed(_row(flags="unconfirmed"))   # confirm 列被裁剪也认
    assert not rlc_report.unconfirmed(_row(verdict="undecided", confirm=""))
    assert not rlc_report.unconfirmed(_row(confirm="confirmed"))


def test_row_notes_give_each_recheck_outcome_its_own_note():
    rows = [_row(confirm="timer_mismatch"),
            _row(row=1, confirm="insufficient_rounds"),
            _row(row=2, confirm="not_reproduced")]
    assert rlc_report.row_notes(rows, None, {}, "main:kron@kron") == [
        "复测计时口径不一致，主轮次结论未确认（1 行，未计入汇总）",
        "复测样本不足，主轮次结论未确认（1 行，未计入汇总）",
        "复测未重现（1 行，未计入汇总）"]


def test_row_notes_fall_back_to_a_generic_note_without_a_confirm_column():
    """A --csv whose `confirm` column was dropped keeps only `flags=unconfirmed`. The row
    is still excluded, so the note has to say why instead of staying silent — and it must
    not repeat a specific reason when one is present."""
    rows = [_row(flags="unconfirmed", confirm="", gap=0.2)]
    assert rlc_report.row_notes(rows, None, {}, "main:kron@kron") == [
        "结果未确认（1 行，未计入汇总）"]
    assert rlc_report.row_notes(rows, None, {}, "main:kron@kron", dropping=False) == [
        "结果未确认（1 行，仍计入汇总）"]
    named = [_row(flags="unconfirmed", confirm="not_reproduced", gap=0.2)]
    assert rlc_report.row_notes(named, None, {}, "main:kron@kron") == [
        "复测未重现（1 行，未计入汇总）"]
    mismatch = [_row(flags="unconfirmed;timer_mismatch", gap=0.2)]
    assert rlc_report.row_notes(mismatch, None, {}, "main:kron@kron") == [
        "A/B 计时口径不一致，性能差异不可直接比较（未计入收益/损失汇总，1 行）"]


def test_summarizing_reports_the_generic_note_when_confirm_is_missing(tmp_path):
    """Same path end to end through the CSV reader: no confirm column, only flags."""
    path = tmp_path / "rows.csv"
    path.write_text(
        "op,run,op_name,dtype,row,gap,median_a_ms,median_b_ms,gate,flags,verdict\n"
        "main:kron@kron,main/kron,kron,torch.bfloat16,0,0.5,0.05,0.02,changed,unconfirmed,"
        "faster\n")
    row = rlc_report.read_ab_csv(path)[0]
    assert rlc_report.unconfirmed(row)
    stats = rlc_report.summarizing([row], {}, None, {}, "main:kron@kron", "bf16")
    assert stats["rows"] == 1
    assert stats["gain"] == (None, None, 0)
    assert "结果未确认（1 行，未计入汇总）" in stats["note"]


def test_summarizing_excludes_not_reproduced_and_short_rechecks():
    rows = [_row(gap=0.5, confirm="not_reproduced"),
            _row(row=1, gap=0.4, confirm="insufficient_rounds"),
            _row(row=2, gap=0.2)]
    stats = rlc_report.summarizing(rows, {}, None, {}, "main:kron@kron", "bf16")
    assert stats["rows"] == 3
    assert stats["gain"] == (0.2, 0.2, 1)
    assert "复测未重现（1 行，未计入汇总）" in stats["note"]
    assert "复测样本不足，主轮次结论未确认（1 行，未计入汇总）" in stats["note"]


def test_summarizing_counts_unconfirmed_rows_when_asked():
    """--include-unconfirmed puts them back, and the note says the numbers contain them."""
    rows = [_row(gap=0.5, flags="timer_mismatch"),
            _row(row=1, gap=-0.2, a=0.02, b=0.025, confirm="timer_mismatch")]
    stats = rlc_report.summarizing(rows, {}, None, {}, "main:kron@kron", "bf16",
                                   include_unconfirmed=True)
    assert stats["rows"] == 2
    assert stats["gain"] == (0.5, 0.5, 1)
    assert stats["loss"] == (-0.2, -0.2, 1)
    assert "仍计入汇总" in stats["note"]


def test_op_entry_saving_and_cost_skip_unconfirmed_rows():
    """The op's saving/cost maxima are a second summary path and follow the same rule."""
    rows = [_row(gap=0.5, a=0.05, b=0.02),
            _row(row=1, gap=0.2, a=0.10, b=0.02, flags="timer_mismatch")]
    groups = {g.dtype: g.rows for g in rlc_report.build_groups(rows)}
    entry = rlc_report.op_entry("main:kron@kron", rows, groups, {}, None, {}, {}, {})
    assert entry["rows"] == 2
    assert entry["saving"] == pytest.approx(0.03)          # 只算干净行 0.05-0.02
    entry = rlc_report.op_entry("main:kron@kron", rows, groups, {}, None, {}, {}, {},
                                include_unconfirmed=True)
    assert entry["saving"] == pytest.approx(0.08)          # 0.10-0.02 回来了


def test_summarizing_reads_the_gate_convert_layout_count(tmp_path):
    (tmp_path / "base" / "H1").mkdir(parents=True)
    (tmp_path / "new" / "H1").mkdir(parents=True)
    (tmp_path / "base" / "H1" / "k.digest.json").write_text('{"ops": {"ttg.convert_layout": 2}}')
    (tmp_path / "new" / "H1" / "k.digest.json").write_text('{"ops": {"ttg.convert_layout": 0}}')
    gate = {("main/kron", "kron", "torch.bfloat16", 0): ["H1"]}
    stats = rlc_report.summarizing([_row()], gate, tmp_path, {}, "main:kron@kron", "bf16")
    assert stats["cvt"] == (2, 0, False)
    assert rlc_report.fmt_cvt(stats["cvt"]) == "2→0"


def test_summarizing_marks_an_unchanged_kernel_with_its_convert_count(tmp_path):
    (tmp_path / "base" / "H1").mkdir(parents=True)
    (tmp_path / "new" / "H1").mkdir(parents=True)
    (tmp_path / "base" / "H1" / "k.digest.json").write_text('{"ops": {"ttg.convert_layout": 2}}')
    (tmp_path / "new" / "H1" / "k.digest.json").write_text('{"ops": {"ttg.convert_layout": 2}}')
    kernels = {("main/kron", "kron", "torch.bfloat16", 0): ["H1"]}
    stats = rlc_report.summarizing([_row()], {}, tmp_path, {}, "main:kron@kron", "bf16",
                                   None, kernels)
    assert stats["cvt"] == (2, 2, True)
    assert rlc_report.fmt_cvt(stats["cvt"]) == "2 (kernel unchanged)"


def test_summarizing_reads_a_single_op_gate_under_its_empty_run(tmp_path):
    """ir_gate.py records a single-op run dir with run "" (see RunDir.run), and
    ab_compare.gate_flag falls back to that key. The report must too: otherwise a
    changed kernel reads as "kernel unchanged" and the convert count is lost."""
    (tmp_path / "base" / "H1").mkdir(parents=True)
    (tmp_path / "new" / "H1").mkdir(parents=True)
    (tmp_path / "base" / "H1" / "k.digest.json").write_text('{"ops": {"ttg.convert_layout": 5}}')
    (tmp_path / "new" / "H1" / "k.digest.json").write_text('{"ops": {"ttg.convert_layout": 2}}')
    gate = {("", "kron", "torch.bfloat16", 0): ["H1"]}
    kernels = {("main/kron", "kron", "torch.bfloat16", 0): ["H1"]}
    stats = rlc_report.summarizing([_row()], gate, tmp_path, {}, "main:kron@kron", "bf16",
                                   None, kernels)
    assert rlc_report.fmt_cvt(stats["cvt"]) == "5→2"


def test_row_entry_falls_back_to_the_empty_run():
    table = {("", "kron", "torch.bfloat16", 0): ["H1"]}
    assert rlc_report.row_entry(table, "main/kron", "kron", "torch.bfloat16", 0) == ["H1"]
    assert rlc_report.row_entry(table, "main/kron", "kron", "torch.bfloat16", 9) == []
    assert rlc_report.row_entry(None, "main/kron", "kron", "torch.bfloat16", 0) == []


def test_main_reads_a_single_op_gate(monkeypatch, tmp_path):
    """The whole report path, with the gate of a single-op run dir: the changed
    kernel and its convert count must survive the empty run key."""
    gate_dir = tmp_path / "gate"
    for side, count in (("base", 5), ("new", 2)):
        (gate_dir / side / "H1").mkdir(parents=True)
        (gate_dir / side / "H1" / "k.digest.json").write_text(
            '{"ops": {"ttg.convert_layout": %d}}' % count)
    (gate_dir / "rows.tsv").write_text(
        "run\top\top_name\tdtype\trow\tshape\tchanged\tchanged_kernels\n"
        "\tmain:kron@kron\tkron\ttorch.bfloat16\t0\t4x4\t1\tH1\n", encoding="utf-8")
    csv_path = tmp_path / "rows.csv"
    csv_path.write_text(
        "op,run,op_name,dtype,row,gap,median_a_ms,median_b_ms,gate,flags,verdict\n"
        "main:kron@kron,main/kron,kron,torch.bfloat16,0,0.2,1.0,0.8,changed,,faster\n",
        encoding="utf-8")
    seen = {}
    monkeypatch.setattr(rlc_report, "write_xlsx",
                        lambda path, rows, title, env: seen.update(rows=rows))
    monkeypatch.setattr(sys, "argv", ["rlc_report.py", "--csv", str(csv_path),
                                      "--gate", str(gate_dir),
                                      "--out", str(tmp_path / "out.xlsx")])
    rlc_report.main()
    assert rlc_report.fmt_cvt(seen["rows"][0]["groups"]["bf16"]["cvt"]) == "5→2"
    path = tmp_path / "r1_a" / "main" / "kron" / "rows.tsv"
    path.parent.mkdir(parents=True)
    path.write_text("op_name\tdtype\trow\tshape\tkernels\n"
                    "kron\ttorch.bfloat16\t0\t4x4\tH1\n"
                    "kron\ttorch.bfloat16\t1\t8x8\tH1,H2\n", encoding="utf-8")
    assert rlc_report.load_run_kernels(tmp_path) == {
        ("main/kron", "kron", "torch.bfloat16", 0): ["H1"],
        ("main/kron", "kron", "torch.bfloat16", 1): ["H1", "H2"],
    }
    assert rlc_report.load_run_kernels(None) == {}


def test_dtype_groups_and_formatters():
    assert rlc_report.dtype_group("torch.bfloat16") == "bf16"
    assert rlc_report.dtype_group("torch.float16") == "fp16"
    assert rlc_report.dtype_group("torch.complex64") == "fp8 / complex64"
    assert rlc_report.dtype_group("torch.int32") is None
    assert rlc_report.fmt_time(0.0) == "<0.1 µs"
    assert rlc_report.fmt_time(0.0005) == "0.5 µs"
    assert rlc_report.fmt_time(0.05) == "50 µs"
    assert rlc_report.fmt_time(12.09) == "12.09 ms"


def test_op_entry_counts_the_dtypes_the_table_has_no_column_for():
    """A dtype without a column (float32, ...) is not a dtype to drop: its rows still
    name a gain or a regression of the op, and the note says they are not shown."""
    rows = [_row(gap=0.05), _row(gap=-0.5, row=1, dtype="torch.float32", a=2.0, b=3.0)]
    groups = {g.dtype: g.rows for g in rlc_report.build_groups(rows)}
    entry = rlc_report.op_entry("main:kron@kron", rows, groups, {}, None, {}, {}, {})
    assert entry["rows"] == 2
    assert list(entry["groups"]) == ["bf16"]
    assert entry["cost"] == pytest.approx(1.0)   # the float32 row's regression
    assert "float32 1 行" in entry["note"]


def test_op_entry_keeps_the_warnings_of_an_ungrouped_dtype():
    """A row the table has no column for still carries do_bench timing, a replay/
    config mismatch and the manual note — dropping them would make a result that no
    longer uses the same config look fine."""
    rows = [_row(dtype="torch.float32", flags="do_bench;replay_fallback", gate="unchanged",
                 verdict="slower")]
    groups = {g.dtype: g.rows for g in rlc_report.build_groups(rows)}
    entry = rlc_report.op_entry("main:kron@kron", rows, groups, {}, None,
                                {"main:kron@kron": "手工：选到不同 autotune"},
                                {"main/kron": "host 同步"}, {})
    assert "do_bench 计时（表级标记覆盖 1/1 行；host 同步）" in entry["note"]
    assert "replay 回退，config 不一致" in entry["note"]
    assert "手工：选到不同 autotune" in entry["note"]
    assert "float32 1 行" in entry["note"]


def test_op_entry_gathers_warnings_from_every_dtype():
    """The op's note is not the union of its dtype groups' notes: a warning only the
    float32 row carries still has to appear next to the bf16 columns."""
    rows = [_row(gap=0.05), _row(dtype="torch.float32", row=1, gap=-0.1,
                                 flags="replay_fallback", verdict="slower")]
    groups = {g.dtype: g.rows for g in rlc_report.build_groups(rows)}
    entry = rlc_report.op_entry("main:kron@kron", rows, groups, {}, None, {}, {}, {})
    assert "replay 回退，config 不一致" in entry["note"]
    assert entry["groups"]["bf16"]["note"] == ""


def test_op_entry_labels_a_group_only_warning_with_its_dtype():
    """The table has one note per op, so a warning that only one dtype's rows produce
    has to carry that dtype: bf16 keeping its kernels is a bf16-level statement and
    stops being true for the op once the fp16 rows changed."""
    rows = [_row(gate="unchanged", verdict="slower", gap=-0.1, row=0),
            _row(dtype="torch.float16", gate="changed", verdict="faster", gap=0.2, row=1)]
    groups = {g.dtype: g.rows for g in rlc_report.build_groups(rows)}
    entry = rlc_report.op_entry("main:kron@kron", rows, groups, {}, None, {}, {}, {})
    assert entry["note"] == "bf16 指令未增删"
    assert entry["groups"]["bf16"]["note"] == "指令未增删"   # the group keeps it unlabelled


def test_op_entry_keeps_the_plain_unchanged_warning_when_every_dtype_agrees():
    rows = [_row(gate="unchanged", verdict="slower", gap=-0.1)]
    groups = {g.dtype: g.rows for g in rlc_report.build_groups(rows)}
    entry = rlc_report.op_entry("main:kron@kron", rows, groups, {}, None, {}, {}, {})
    assert entry["note"] == "指令未增删"


def test_ungrouped_note_is_empty_when_every_row_has_a_column():
    assert rlc_report.ungrouped_note([_row(dtype="torch.bfloat16"),
                                      _row(dtype="torch.float8_e4m3fn")]) == ""


def test_main_reports_an_op_that_only_ran_float32(monkeypatch, tmp_path):
    """The CSV is the source of the op list: an op whose rows are all float32 still
    gets its table row (and its saving/cost), instead of the report coming back
    empty while the input had rows."""
    csv_path = tmp_path / "rows.csv"
    csv_path.write_text(
        "op,run,op_name,dtype,row,gap,median_a_ms,median_b_ms,gate,flags,verdict\n"
        "main:kron@kron,main/kron,kron,torch.float32,0,-0.2,1.0,1.2,changed,,slower\n",
        encoding="utf-8")
    seen = {}
    monkeypatch.setattr(rlc_report, "write_xlsx",
                        lambda path, rows, title, env: seen.update(rows=rows))
    monkeypatch.setattr(sys, "argv", ["rlc_report.py", "--csv", str(csv_path),
                                      "--out", str(tmp_path / "out.xlsx")])
    rlc_report.main()
    assert len(seen["rows"]) == 1
    entry = seen["rows"][0]
    assert entry["label"] == "main:kron@kron" and entry["rows"] == 1
    assert entry["groups"] == {} and entry["cost"] == pytest.approx(0.2)
    assert "float32 1 行" in entry["note"]


def test_notes_are_automatic_and_manual(monkeypatch, tmp_path):
    rows = [_row(gap=-0.2, flags="do_bench", verdict="slower", gate="unchanged")]
    notes = {"main:kron@kron": "选到不同 autotune"}
    stats = rlc_report.summarizing(rows, {}, None, notes, "main:kron@kron", "bf16")
    assert "do_bench" in stats["note"]
    assert "选到不同 autotune" in stats["note"]


def test_xlsx_has_the_expected_groups(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    path = tmp_path / "out.xlsx"
    rows = [{"name": "kron", "label": "main:kron@kron", "rows": 2,
             "groups": {"bf16": rlc_report.summarizing(
                 [_row(gap=1.5, row=0), _row(gap=-0.1, row=1)], {}, None, {},
                 "main:kron@kron", "bf16")},
             "saving": 0.03, "cost": 0.001, "note": ""}]
    rlc_report.write_xlsx(path, rows, "title", "env")
    ws = load_workbook(path).active
    assert ws.max_column == 27
    assert ws.cell(4, 3).value == "bf16"
    assert ws.cell(4, 10).value == "fp16"
    assert ws.cell(4, 17).value == "fp8 / complex64"
    assert ws.cell(6, 3).value == "最弱"
    assert ws.cell(7, 2).value == "kron"
    assert ws.cell(7, 3).value == "+150.0%"
    assert ws.cell(7, 6).value == "-10.0%"


def _one_op_table(tmp_path, cell):
    path = tmp_path / "one.csv"
    path.write_text(f"#,算子\n1,\"{cell}\"\n", encoding="utf-8")
    return path


def test_resolve_runs_both_repos_when_both_have_the_op(monkeypatch, tmp_path):
    inv = {"main": [{"op": "outer", "op_file": "outer", "collected": "yes"}],
           "vllm": [{"op": "outer", "op_file": "outer", "collected": "yes"}]}
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "outer")))
    assert specs == ["main:outer@outer", "vllm:outer@outer"]
    assert [r["status"] for r in report] == ["both", "both"]
    assert not errors


def test_resolve_uses_vllm_for_an_op_that_only_exists_there(monkeypatch, tmp_path):
    inv = {"main": [{"op": "kron", "op_file": "kron", "collected": "yes"}],
           "vllm": [{"op": "fused_marlin_moe_w8a16_fp8",
                     "op_file": "fused_marlin_moe_w8a16_fp8", "collected": "yes"}]}
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "fused_marlin_moe_w8a16_fp8")))
    assert specs == ["vllm:fused_marlin_moe_w8a16_fp8@fused_marlin_moe_w8a16_fp8"]
    assert report[0]["status"] == "vllm"
    assert not errors


def test_resolve_records_an_op_missing_from_both_repos(monkeypatch, tmp_path):
    inv = {"main": [{"op": "kron", "op_file": "kron", "collected": "yes"}],
           "vllm": [{"op": "outer", "op_file": "outer", "collected": "yes"}]}
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "no_such_op")))
    assert specs == []
    assert report == [{"line": 2, "display": "no_such_op", "repo": "",
                       "op": "no_such_op", "op_file": "", "status": "missing"}]
    assert not errors  # recorded, not fatal


def test_resolve_pins_one_op_file_from_the_first_line(monkeypatch, tmp_path):
    """`main:kron@kron` on a single line is the same spec as the two-line form:
    it must pin that benchmark file instead of being read as an op named
    `kron@kron`, and `kron@kron` must pin it in both repos."""
    inv = {"main": [{"op": "kron", "op_file": "kron", "collected": "yes"},
                    {"op": "kron", "op_file": "kron_other", "collected": "yes"}],
           "vllm": [{"op": "kron", "op_file": "kron", "collected": "yes"}]}
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)

    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "main:kron@kron")))
    assert specs == ["main:kron@kron"]
    assert [r["status"] for r in report] == ["main"]
    assert not errors

    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "kron\nmain:kron@kron")))
    assert specs == ["main:kron@kron"]
    assert not errors

    # no repo prefix: the pinned file is looked up in both repos
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "kron@kron")))
    assert specs == ["main:kron@kron", "vllm:kron@kron"]
    assert [r["status"] for r in report] == ["both", "both"]
    assert not errors


def test_resolve_reports_a_pinned_op_file_that_is_missing(monkeypatch, tmp_path):
    inv = {"main": [{"op": "kron", "op_file": "kron", "collected": "yes"}]}
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "main:kron@nope")))
    assert specs == []
    assert errors == ["line 2: no main:kron@nope in the inventory"]
    assert report == [{"line": 2, "display": "main:kron@nope", "repo": "main",
                       "op": "kron", "op_file": "nope", "status": "missing"}]


def test_resolve_matches_the_benchmark_op_name(monkeypatch, tmp_path):
    inv = {"main": [{"op": "perf_fft", "op_file": "fft", "op_names": "fft",
                     "collected": "yes"}]}
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "fft")))
    assert specs == ["main:perf_fft@fft"]
    assert report[0]["status"] == "main"
    assert not errors


def test_resolve_matches_the_op_file_when_op_names_is_empty(monkeypatch, tmp_path):
    inv = {"main": [{"op": "scaled_grouped_mm_benchmark", "op_file": "scaled_grouped_mm",
                     "op_names": "", "collected": "yes"}]}
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "scaled_grouped_mm")))
    assert specs == ["main:scaled_grouped_mm_benchmark@scaled_grouped_mm"]
    assert not errors


def test_resolve_expands_several_tests_with_one_op_name(monkeypatch, tmp_path):
    inv = {"main": [{"op": "fused_moe_fp8", "op_file": "fused_moe_fp8",
                     "op_names": "fused_experts_impl", "collected": "yes"},
                    {"op": "fused_moe_int8", "op_file": "fused_moe_int8",
                     "op_names": "fused_experts_impl", "collected": "yes"}]}
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "fused_experts_impl")))
    assert specs == ["main:fused_moe_fp8@fused_moe_fp8",
                     "main:fused_moe_int8@fused_moe_int8"]
    assert [r["status"] for r in report] == ["multi", "multi"]
    assert not errors


def test_resolve_expands_all_tests_in_both_repos(monkeypatch, tmp_path):
    inv = {"main": [{"op": "perf_router_gemm", "op_file": "blas_perf_parallel",
                     "op_names": "router_gemm", "collected": "yes"},
                    {"op": "perf_router_gemm", "op_file": "router_gemm",
                     "op_names": "router_gemm", "collected": "yes"}],
           "vllm": [{"op": "perf_router_gemm", "op_file": "router_gemm",
                     "op_names": "router_gemm", "collected": "yes"}]}
    monkeypatch.setattr(rlc_ops.run_batch, "load_inventory", lambda: inv)
    specs, report, errors = rlc_ops.resolve(rlc_ops.read_table(
        _one_op_table(tmp_path, "router_gemm")))
    assert specs == ["main:perf_router_gemm@blas_perf_parallel",
                     "main:perf_router_gemm@router_gemm",
                     "vllm:perf_router_gemm@router_gemm"]
    assert [r["status"] for r in report] == ["multi", "multi", "multi"]
    assert not errors


def test_display_name_marks_multi_and_both(tmp_path):
    sources = {
        "main:fused_moe_fp8@fused_moe_fp8": {"display": "fused_experts_impl",
                                             "status": "multi"},
        "vllm:outer@outer": {"display": "outer", "status": "both"},
    }
    assert rlc_report.display_name("main:fused_moe_fp8@fused_moe_fp8", {}, sources) == \
        "fused_moe_fp8\nfused_experts_impl · main · multi"
    assert rlc_report.display_name("vllm:outer@outer", {}, sources) == \
        "outer\nouter · vllm · both"
    # without a source row the old same-name rule applies
    assert rlc_report.display_name("main:outer@outer", {"outer": 2}, {}) == \
        "outer\nmain:outer@outer"
    assert rlc_report.display_name("main:kron@kron", {"kron": 1}, {}) == "kron"


def test_display_name_keeps_the_op_file_for_one_name_in_several_files():
    """Two benchmark files with one op name (the `multi` case) must stay apart:
    the second line names the table row both were expanded from, so without the
    op_file on the first line the two rows read exactly the same."""
    sources = {
        "main:perf_router_gemm@blas_perf_parallel": {"display": "router_gemm",
                                                     "status": "multi"},
        "main:perf_router_gemm@router_gemm": {"display": "router_gemm", "status": "multi"},
    }
    op_files = {("main", "perf_router_gemm"): {"blas_perf_parallel", "router_gemm"}}
    a = rlc_report.display_name("main:perf_router_gemm@blas_perf_parallel", {}, sources, op_files)
    b = rlc_report.display_name("main:perf_router_gemm@router_gemm", {}, sources, op_files)
    assert a == "perf_router_gemm@blas_perf_parallel\nrouter_gemm · main · multi"
    assert b == "perf_router_gemm@router_gemm\nrouter_gemm · main · multi"
    assert a != b
    # one file per op name keeps the plain first line
    assert rlc_report.display_name("main:kron@kron", {}, {}, {("main", "kron"): {"kron"}}) == "kron"
    assert rlc_report.display_name("main:kron@kron", {"kron": 2}, {},
                                   {("main", "kron"): {"kron"}}) == "kron\nmain:kron@kron"


def test_load_sources_reads_the_provenance_csv(tmp_path):
    path = tmp_path / "sources.csv"
    path.write_text("line,display,repo,op,op_file,status\n"
                    "7,router_gemm,main,perf_router_gemm,router_gemm,multi\n",
                    encoding="utf-8")
    sources = rlc_report.load_sources(path)
    assert sources["main:perf_router_gemm@router_gemm"]["display"] == "router_gemm"
    assert sources["main:perf_router_gemm@router_gemm"]["status"] == "multi"
    assert rlc_report.load_sources(None) == {}
    assert rlc_report.load_sources(tmp_path / "missing.csv") == {}


def test_read_table_order_unpacks_the_resolve_result(monkeypatch, tmp_path):
    import rlc_ops
    monkeypatch.setattr(rlc_ops, "resolve",
                        lambda entries: (["main:kron@kron"], [], []))
    table = tmp_path / "t.csv"
    table.write_text("#,算子\n1,kron\n", encoding="utf-8")
    assert rlc_report.read_table_order(table, ["main:kron@kron"]) == \
        {"main:kron@kron": 0}


def test_load_notes_ignores_a_missing_file(tmp_path):
    assert rlc_report.load_notes(None) == {}
    assert rlc_report.load_notes(tmp_path / "missing.csv") == {}


def test_load_do_bench_reasons_reads_the_marker(tmp_path):
    log = tmp_path / "r1_a" / "main" / "kron" / "run.log"
    log.parent.mkdir(parents=True)
    log.write_text(
        "BENCHMARK_DIRECT_NO_CUDAGRAPH op=kron phase=host_sync reason=_HostSync "
        "detail='tensor.item()' \n"
        "BENCHMARK_DIRECT_NO_CUDAGRAPH op=kron phase=warmup reason=OutOfResources "
        "detail='out of memory'\n", encoding="utf-8")
    assert rlc_report.load_do_bench_reasons(tmp_path) == {
        "main/kron": "host 同步；预热显存不足"}
    assert rlc_report.load_do_bench_reasons(None) == {}
    assert rlc_report.load_do_bench_reasons(tmp_path / "missing") == {}


def test_describe_do_bench_covers_the_phases():
    assert rlc_report.describe_do_bench("host_sync", "_HostSync") == "host 同步"
    assert rlc_report.describe_do_bench("capability", "GraphUnsupported") == "平台无 graph API"
    assert rlc_report.describe_do_bench("warmup", "OutOfResources") == "预热显存不足"
    assert rlc_report.describe_do_bench("capture", "AcceleratorError") == \
        "capture 失败: AcceleratorError"


def test_summarizing_appends_the_do_bench_reason():
    rows = [_row(gap=-0.2, flags="do_bench", verdict="slower")]
    stats = rlc_report.summarizing(rows, {}, None, {}, "main:kron@kron", "bf16",
                                   {"main/kron": "host 同步"})
    assert "do_bench 计时（表级标记覆盖 1/1 行；host 同步）" in stats["note"]


def test_do_bench_note_counts_the_rows_each_way():
    """Rows of one (op, dtype) table all carry the flag once the table is marked, so a
    partial count only comes from a hand-built or historical CSV. The note still names
    both counts, and says 表级标记 so the number is not read as a fallback ratio."""
    rows = [_row(gap=0.1, row=0, flags="do_bench"),
            _row(gap=0.2, row=1),
            _row(gap=0.3, row=2)]
    stats = rlc_report.summarizing(rows, {}, None, {}, "main:kron@kron", "bf16",
                                   {"main/kron": "host 同步"})
    assert "do_bench 计时（表级标记覆盖 1/3 行；host 同步）" in stats["note"]
    rows = [_row(gap=0.1, row=0, flags="do_bench"), _row(gap=0.2, row=1, flags="do_bench")]
    stats = rlc_report.summarizing(rows, {}, None, {}, "main:kron@kron", "bf16")
    assert "do_bench 计时（表级标记覆盖 2/2 行）" in stats["note"]


def test_one_fallback_marker_marks_the_whole_table_end_to_end(tmp_path):
    """The note is a table-level warning, and this is the path that shows why: a run.log
    with ONE fallback marker and an 8-row table comes back with all 8 rows flagged, so
    the note reads 8/8 — a reader must not take that as "8 shapes fell back"."""
    header = "Operator: kron  (dtype=torch.bfloat16, mode=kernel, level=core)\n"
    body = "".join(f"SUCCESS  0.010000  {0.1 + i / 100:.6f}  1.000  [torch.Size([{i + 1}])]\n"
                   for i in range(8))
    marker = ("BENCHMARK_DIRECT_NO_CUDAGRAPH op=kron phase=host_sync reason=_HostSync "
              "detail='tensor.item()'\n")
    for side, prefix in (("a", marker), ("b", "")):
        run = tmp_path / f"r1_{side}" / "main" / "kron"
        run.mkdir(parents=True)
        (run / "run.log").write_text(prefix + header + body, encoding="utf-8")

    samples = ab_compare.load_samples(tmp_path)
    assert len(samples) == 8
    assert all(s.nocg for s in samples)          # one marker, eight flagged rows
    rows = [rlc_report.Row("main:kron@kron", s.run, s.op_name, s.dtype, s.row, 0.0,
                           s.a[0], s.b[0], "changed",
                           "do_bench" if s.nocg else "", "same")
            for s in samples]
    stats = rlc_report.summarizing(rows, {}, None, {}, "main:kron@kron", "bf16",
                                   {"main/kron": "host 同步"})
    assert "do_bench 计时（表级标记覆盖 8/8 行；host 同步）" in stats["note"]


def test_do_bench_note_counts_rows_behind_marked_tables():
    """A fraction below 1 means the note spans several tables, not that some rows fell
    back: here the bf16 table is marked (2 rows) and the fp16 one is not (2 rows), so the
    op's note reads 2/4 in the marked-tables sense."""
    rows = [_row(row=0, flags="do_bench"), _row(row=1, flags="do_bench"),
            _row(row=2, dtype="torch.float16"), _row(row=3, dtype="torch.float16")]
    groups = {g.dtype: g.rows for g in rlc_report.build_groups(rows)}
    entry = rlc_report.op_entry("main:kron@kron", rows, groups, {}, None, {}, {}, {})
    assert "do_bench 计时（表级标记覆盖 2/4 行）" in entry["note"]
